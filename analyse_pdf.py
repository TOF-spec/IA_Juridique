import os
import re

from mistralai.client import Mistral
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


# ============================================================
# CONFIGURATION
# ============================================================

API_KEY = os.environ["MISTRAL_API_KEY"]

AGENT_ID = "ag_01a0a05ef1347680b9297f3dff763223"

PDF_PATH = "contrat_test_services.pdf"

WORD_PATH = "rapport_analyse_aurore.docx"


# ============================================================
# CONNEXION À MISTRAL
# ============================================================

client = Mistral(
    api_key=API_KEY
)


# ============================================================
# 1. UPLOAD DU PDF
# ============================================================

print()
print("==========================================")
print("       ANALYSE DU CONTRAT")
print("==========================================")
print()

print("1. Envoi du PDF à Mistral...")

with open(PDF_PATH, "rb") as pdf_file:

    uploaded_file = client.files.upload(
        file={
            "file_name": PDF_PATH,
            "content": pdf_file,
        },
        purpose="ocr",
    )

print("PDF envoyé.")
print("ID du fichier :", uploaded_file.id)


# ============================================================
# 2. OCR DU PDF
# ============================================================

print()
print("2. Extraction du texte du contrat...")

ocr_response = client.ocr.process(
    model="mistral-ocr-latest",
    document={
        "type": "file",
        "file_id": uploaded_file.id,
    },
)

print("OCR terminé.")


# ============================================================
# 3. EXTRACTION DU TEXTE
# ============================================================

texte_contrat = ""

for page in ocr_response.pages:

    if hasattr(page, "markdown"):

        texte_contrat += page.markdown
        texte_contrat += "\n\n"


if not texte_contrat.strip():

    raise Exception(
        "Aucun texte n'a été extrait du PDF."
    )


print("Texte extrait :", len(texte_contrat), "caractères")


# ============================================================
# 4. DEMANDE D'ANALYSE À AURORE
# ============================================================

question = f"""
Tu es Aurore, assistant juridique interne chargé de la revue
de contrats.

Tu dois analyser le contrat ci-dessous en utilisant le
Référentiel juridique interne disponible dans ta bibliothèque.

============================================================
CONTRAT À ANALYSER
============================================================

{texte_contrat}

============================================================
RÈGLES DE QUALITÉ
============================================================

Tu dois distinguer strictement quatre niveaux d'information :

A. TEXTE DU CONTRAT

Indique uniquement ce qui est effectivement écrit dans le contrat.

B. RÉFÉRENTIEL INTERNE

Indique uniquement les règles qui figurent réellement dans
le Référentiel juridique interne.

C. ANALYSE

Présente ton raisonnement sur le contrat.

D. RECOMMANDATION

Présente les actions ou modifications que tu recommandes.

------------------------------------------------------------
RÈGLES STRICTES
------------------------------------------------------------

1. Ne transforme jamais une "cible", une recommandation ou
une préférence interne en obligation juridique.

2. Ne prétends jamais qu'une règle figure dans le référentiel
si elle n'y figure pas explicitement.

3. N'invente aucune information absente du contrat.

4. Si une information nécessaire n'est pas disponible,
indique exactement :

"Information non disponible."

5. Ne présente pas comme certaine une conclusion qui dépend
d'informations manquantes.

6. Lorsque tu proposes une nouvelle rédaction, précise qu'il
s'agit d'une proposition d'Aurore et non d'une clause
provenant du référentiel interne.

7. Pour chaque écart au référentiel, identifie clairement :
- la règle interne ;
- l'article concerné ;
- l'écart constaté ;
- le niveau de risque.

8. Ne confonds jamais :
- risque juridique ;
- écart à la politique interne ;
- point de négociation.

9. Lorsque le référentiel interne ne contient aucune règle
applicable à une clause, indique :

"Aucune règle interne identifiée."

10. Ne fabrique pas de numéro de règle qui n'existe pas
dans le référentiel.

------------------------------------------------------------
NIVEAUX DE RISQUE
------------------------------------------------------------

Utilise uniquement :

CRITIQUE
ÉLEVÉ
MOYEN
FAIBLE

Le niveau de risque doit être justifié.

============================================================
FORMAT DU RAPPORT
============================================================

# 1. SYNTHÈSE EXÉCUTIVE

Présente :

- Parties
- Objet
- Durée
- Montant annuel
- Conditions de paiement
- Droit applicable
- Nombre de points nécessitant une attention particulière

------------------------------------------------------------

# 2. MATRICE DES RISQUES
------------------------------------------------------------

Pour chaque point identifié :

## Article X – [Titre]

**Niveau de risque :**
CRITIQUE / ÉLEVÉ / MOYEN / FAIBLE

**Extrait du contrat :**
"[courte citation exacte du contrat]"

**Constat contractuel :**
Décris uniquement ce qui est écrit dans le contrat.

**Référentiel interne :**
Indique la règle interne applicable.

Si aucune règle ne s'applique :

"Aucune règle interne identifiée."

**Écart au référentiel :**
Décris précisément l'écart.

S'il n'existe aucun écart :

"Aucun écart identifié."

**Analyse :**
Explique le problème ou le risque identifié.

**Certitude :**
FORTE / MOYENNE / À CONFIRMER

**Conséquence pour le Client :**
Explique concrètement la conséquence potentielle.

**Recommandation :**
Indique l'action recommandée.

**Rédaction alternative :**
Propose une rédaction uniquement si cela est pertinent.

------------------------------------------------------------

# 3. ÉCARTS AU RÉFÉRENTIEL INTERNE
------------------------------------------------------------

Présente un tableau avec :

| Article | Règle interne | Écart constaté | Risque |
|---------|---------------|----------------|--------|

------------------------------------------------------------

# 4. CLAUSES MANQUANTES OU INSUFFISAMMENT PRÉCISES
------------------------------------------------------------

Liste les clauses qui semblent manquer ou qui nécessitent
des informations complémentaires.

------------------------------------------------------------

# 5. POINTS À NÉGOCIER
------------------------------------------------------------

Classe les points par priorité :

1. CRITIQUE
2. ÉLEVÉ
3. MOYEN
4. FAIBLE

------------------------------------------------------------

# 6. PROPOSITIONS DE RÉDACTION
------------------------------------------------------------

Regroupe les propositions de nouvelles formulations.

Chaque proposition doit indiquer :

- Article concerné
- Objectif
- Rédaction proposée

Précise qu'il s'agit d'une proposition d'Aurore.

------------------------------------------------------------

# 7. INFORMATIONS MANQUANTES
------------------------------------------------------------

Liste toutes les informations qui empêchent une analyse
complète.

Si aucune information importante ne manque :

"Aucune information critique manquante identifiée."

------------------------------------------------------------

# 8. AVERTISSEMENT
------------------------------------------------------------

Termine obligatoirement par :

"Cette analyse constitue une assistance à la revue contractuelle.
Elle ne remplace pas la validation d'un juriste."
"""


# ============================================================
# 5. ANALYSE PAR AURORE
# ============================================================

print()
print("3. Analyse juridique par Aurore...")
print()

response = client.beta.conversations.start(
    agent_id=AGENT_ID,
    inputs=[
        {
            "role": "user",
            "content": question,
        }
    ],
)


# ============================================================
# 6. RÉCUPÉRATION DU RAPPORT
# ============================================================

rapport = ""

for output in response.outputs:

    if hasattr(output, "content"):

        rapport += output.content


if not rapport.strip():

    raise Exception(
        "Aucun rapport n'a été généré par Aurore."
    )


# ============================================================
# 7. AFFICHAGE DU RAPPORT
# ============================================================

print()
print("==========================================")
print("        RAPPORT D'ANALYSE D'AURORE")
print("==========================================")
print()

print(rapport)

print()
print("==========================================")
print("             FIN DE L'ANALYSE")
print("==========================================")


# ============================================================
# 8. CRÉATION DU DOCUMENT WORD
# ============================================================

print()
print("4. Création du rapport Word...")


document = Document()


# ------------------------------------------------------------
# Style général
# ------------------------------------------------------------

styles = document.styles

styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(10)


# ------------------------------------------------------------
# Titre
# ------------------------------------------------------------

titre = document.add_paragraph()

titre.alignment = WD_ALIGN_PARAGRAPH.CENTER

run = titre.add_run(
    "RAPPORT D'ANALYSE CONTRACTUELLE"
)

run.bold = True
run.font.size = Pt(18)


sous_titre = document.add_paragraph()

sous_titre.alignment = WD_ALIGN_PARAGRAPH.CENTER

run = sous_titre.add_run(
    "Aurore – Assistant juridique interne"
)

run.italic = True
run.font.size = Pt(11)


document.add_paragraph()


# ------------------------------------------------------------
# Informations générales
# ------------------------------------------------------------

p = document.add_paragraph()

run = p.add_run("Contrat analysé : ")
run.bold = True

p.add_run(PDF_PATH)


p = document.add_paragraph()

run = p.add_run("Agent : ")
run.bold = True

p.add_run("Aurore")


document.add_paragraph()


# ------------------------------------------------------------
# Conversion du rapport Markdown en Word
# ------------------------------------------------------------

lignes = rapport.splitlines()

for ligne in lignes:

    ligne = ligne.strip()

    if not ligne:
        continue


    # Titres principaux
    if ligne.startswith("# "):

        texte = ligne[2:].strip()

        document.add_heading(
            texte,
            level=1
        )

        continue


    # Sous-titres
    if ligne.startswith("## "):

        texte = ligne[3:].strip()

        document.add_heading(
            texte,
            level=2
        )

        continue


    # Sous-sous-titres
    if ligne.startswith("### "):

        texte = ligne[4:].strip()

        document.add_heading(
            texte,
            level=3
        )

        continue


    # Listes à puces
    if ligne.startswith("- "):

        texte = ligne[2:].strip()

        document.add_paragraph(
            texte,
            style="List Bullet"
        )

        continue


    # Listes numérotées
    if re.match(r"^\d+\.\s", ligne):

        texte = re.sub(
            r"^\d+\.\s",
            "",
            ligne
        )

        document.add_paragraph(
            texte,
            style="List Number"
        )

        continue


    # Séparateurs Markdown
    if ligne.startswith("---"):

        document.add_paragraph()

        continue


    # Tableaux Markdown
    if ligne.startswith("|"):

        cellules = [
            cellule.strip()
            for cellule in ligne.strip("|").split("|")
        ]

        # Ignorer les lignes de séparation Markdown
        if all(
            set(cellule) <= set("-: ")
            for cellule in cellules
        ):
            continue

        tableau = document.add_table(
            rows=1,
            cols=len(cellules)
        )

        tableau.style = "Table Grid"

        for i, cellule in enumerate(cellules):

            tableau.rows[0].cells[i].text = cellule

        continue


    # Texte normal
    paragraphe = document.add_paragraph()


    # Gestion simple du gras Markdown
    morceaux = re.split(
        r"(\*\*.*?\*\*)",
        ligne
    )

    for morceau in morceaux:

        if morceau.startswith("**") and morceau.endswith("**"):

            run = paragraphe.add_run(
                morceau[2:-2]
            )

            run.bold = True

        else:

            paragraphe.add_run(
                morceau
            )


# ------------------------------------------------------------
# Pied de page
# ------------------------------------------------------------

section = document.sections[0]

footer = section.footer

p = footer.paragraphs[0]

p.alignment = WD_ALIGN_PARAGRAPH.CENTER

p.add_run(
    "Analyse générée par Aurore – Validation juridique humaine requise."
).italic = True


# ------------------------------------------------------------
# Sauvegarde
# ------------------------------------------------------------

document.save(WORD_PATH)


print()
print("==========================================")
print("     RAPPORT WORD CRÉÉ AVEC SUCCÈS")
print("==========================================")
print()
print("Fichier :", WORD_PATH)
print()